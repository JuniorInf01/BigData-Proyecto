"""Conversión TÉCNICA de libros y documentos; reglas analíticas quedan en PySpark.
Debe ejecutarse DESPUÉS de la ingesta verificada a HDFS.
"""
from pathlib import Path
import csv, datetime, json, re, shutil, subprocess, tempfile, zipfile
from xml.etree import ElementTree as ET
from afp_config import raiz_fuentes, STAGING_LOCAL, ruta_hdfs, EVIDENCIAS
from afp_tools import ejecutar, slug, subir_si_no_existe, escribir_jsonl


def libro_a_xlsx(origen, destino_dir):
    """Nunca cambia archivo original; los XLS legacy se convierten a XLSX por LibreOffice."""
    if zipfile.is_zipfile(origen):
        destino=destino_dir/(slug(origen.stem)+'.xlsx')
        shutil.copyfile(origen,destino)
        return destino, 'xlsx_interno_ext_xls'
    if shutil.which('soffice') is None:
        raise RuntimeError('No encuentro soffice (LibreOffice). Solicita que lo instalen en el laboratorio.')
    salida=tempfile.TemporaryDirectory()
    try:
        cmd=['soffice','-env:UserInstallation=file:///tmp/lo_afp_v2_temp','--headless','--convert-to','xlsx','--outdir',salida.name,str(origen)]
        ejecutar(cmd,timeout=240)
        convertidos=list(Path(salida.name).glob('*.xlsx'))
        if len(convertidos)!=1: raise RuntimeError('No se produjo exactamente un archivo XLSX desde '+str(origen))
        destino=destino_dir/(slug(origen.stem)+'.xlsx')
        shutil.copyfile(convertidos[0],destino)
        return destino,'xls_legacy_convertido'
    finally:
        salida.cleanup()


def valor_texto(v):
    if v is None:return ''
    if isinstance(v,datetime.datetime):return v.strftime('%Y-%m-%d')
    if isinstance(v,datetime.date):return v.isoformat()
    return str(v).strip()


def extraer_hojas():
    from openpyxl import load_workbook
    raiz=raiz_fuentes()
    stage=STAGING_LOCAL/'sbs'
    stage.mkdir(parents=True,exist_ok=True)
    libros=[]; hojas=[]
    for origen in sorted((raiz/'Datos Estructurados').iterdir()):
        if origen.suffix.lower() not in ('.xls','.xlsx'):continue
        nombre=slug(origen.stem)
        d=stage/nombre;d.mkdir(parents=True,exist_ok=True)
        ruta, formato=libro_a_xlsx(origen,d)
        libro=load_workbook(ruta,read_only=True,data_only=True)
        libros.append({'archivo_original':origen.name, 'xlsx_convertido':str(ruta), 'formato_detectado':formato, 'hojas':libro.sheetnames})
        for hoja in libro:
            nombre_hoja=slug(hoja.title)
            csv_local=d/(nombre_hoja+'.csv')
            with csv_local.open('w',encoding='utf-8',newline='') as outf:
                wr=csv.writer(outf)
                # 159 columnas en algunas hojas SBS provienen también de formatos heredados.
                # Conservar c0..cN para mantener posiciones de celdas, sin inferir encabezados aquí.
                wr.writerow(['c'+str(i) for i in range(hoja.max_column)])
                for fila in hoja.iter_rows(values_only=True):
                    wr.writerow([valor_texto(x) for x in fila])
            destino=ruta_hdfs('staging','sbs',nombre,nombre_hoja+'.csv')
            subir_si_no_existe(csv_local,destino)
            hojas.append({'libro':nombre,'hoja':hoja.title,'hoja_slug':nombre_hoja,
                           'filas_origen':hoja.max_row,'columnas_origen':hoja.max_column,
                           'hdfs_path':destino})
        libro.close()
    print('Libros SBS:',len(libros),'| hojas separadas:',len(hojas))
    return libros,hojas


def docx_a_texto(path):
    parrafos=[]
    with zipfile.ZipFile(path) as f:
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        nombres=[n for n in f.namelist() if n=='word/document.xml']
        for n in nombres:
            raiz=ET.fromstring(f.read(n))
            for par in raiz.findall('.//w:p',ns):
                fragmentos=[n.text or '' for n in par.findall('.//w:t',ns)]
                if fragmentos:parrafos.append(''.join(fragmentos))
    return '\n'.join(parrafos)


def doc_a_docx(path):
    if shutil.which('soffice') is None:
        raise RuntimeError('LibreOffice necesario para DOC binario: '+str(path))
    temp=tempfile.TemporaryDirectory()
    cmd=['soffice','-env:UserInstallation=file:///tmp/lo_afp_v2_docs','--headless','--convert-to','docx',
         '--outdir',temp.name,str(path)]
    ejecutar(cmd,timeout=180)
    match=list(Path(temp.name).glob('*.docx'))
    if not match:
        temp.cleanup(); raise RuntimeError('Falló conversión DOC: '+str(path))
    return temp,match[0]


def obtener_manifest_memorias(raiz):
    posibles=list((raiz/'Revisión').glob('manifiesto_memorias.csv'))
    if not posibles:return {}
    with posibles[0].open(encoding='utf-8-sig',newline='') as f:
        return {row['ruta'].replace('\\','/'):row for row in csv.DictReader(f)}


def extraer_memorias():
    raiz=raiz_fuentes()
    mapa=obtener_manifest_memorias(raiz)
    directorio=raiz/'Datos No estructurados'
    stage=STAGING_LOCAL/'documentos'
    stage.mkdir(parents=True,exist_ok=True)
    documentos=[]; errores=[]
    for origen in sorted(directorio.rglob('*')):
        if not origen.is_file():continue
        relativo=str(origen.relative_to(raiz)).replace('\\','/')
        meta=mapa.get(relativo,{})
        anio=meta.get('anio_ejercicio')
        if not anio:
            m=re.search(r'20(?:1[5-9]|2[0-5])',origen.name)
            anio=m.group() if m else None
        if not anio:
            errores.append({'archivo':relativo,'error':'Sin año en nombre ni manifiesto; no se inventa'})
            continue
        afp=origen.parent.name.replace('Memoria ','').strip()
        salida=stage/slug(afp)/(slug(origen.stem)+'.txt')
        salida.parent.mkdir(parents=True,exist_ok=True)
        try:
            formato=origen.suffix.lower()
            if formato=='.pdf':
                if not shutil.which('pdftotext'): raise RuntimeError('Falta pdftotext')
                cmd=['pdftotext','-layout',str(origen),str(salida)]
                ejecutar(cmd,timeout=240)
                texto=salida.read_text(encoding='utf-8',errors='replace')
            elif formato=='.docx':
                texto=docx_a_texto(origen)
                salida.write_text(texto,encoding='utf-8')
            elif formato=='.doc':
                temp,convertido=doc_a_docx(origen)
                try: texto=docx_a_texto(convertido)
                finally:temp.cleanup()
                salida.write_text(texto,encoding='utf-8')
            else:raise RuntimeError('Extensión desconocida: '+formato)
            if len(texto.strip())<100:
                raise RuntimeError('Texto insuficiente: revisar si PDF escaneado, sin aplicar OCR automáticamente')
            destino=ruta_hdfs('staging','documentos',slug(afp),salida.name)
            subir_si_no_existe(salida,destino)
            documentos.append({'afp':afp,'anio_ejercicio':int(anio),
                'parte':int(meta.get('parte',1)),
                'disponible_desde_regla':meta.get('disponible_desde_regla',f'{int(anio)+1}-04-01'),
                'archivo_original':relativo,'hdfs_path':destino,'caracteres_extraidos':len(texto)})
        except Exception as exc:
            errores.append({'archivo':relativo,'error':str(exc)})
    print('Documentos extraídos:',len(documentos),'| errores:',len(errores))
    if errores:print('ERRORES (no inventar contenido):',json.dumps(errores[:10],ensure_ascii=False,indent=2))
    return documentos,errores


def ejecutar_preparacion():
    EVIDENCIAS.mkdir(parents=True,exist_ok=True)
    libros,hojas=extraer_hojas()
    documentos,errores=extraer_memorias()
    registro={'libros':libros,'hojas':hojas,'documentos':documentos,'errores':errores}
    local=EVIDENCIAS/'03_staging_registro.json'
    local.write_text(json.dumps(registro,ensure_ascii=False,indent=2),encoding='utf-8')
    if len(libros)!=3:print('ATENCIÓN: se esperaban 3 libros SBS.')
    if len(documentos)!=48:print('ATENCIÓN: hay documentos no extraídos: revisar el registro antes de modelar.')
    if len(errores):raise RuntimeError('Extracción parcial; resolver errores del registro antes de continuar')
    return registro
